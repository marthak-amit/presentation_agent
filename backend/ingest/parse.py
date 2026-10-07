"""PPTX -> structured slide text (title, body, tables, speaker notes)."""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass
class ParsedSlide:
    n: int
    title: str = ""
    body: str = ""
    tables: str = ""
    notes: str = ""
    body_lines: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        d = asdict(self)
        d.pop("body_lines")
        return d

    def index_text(self) -> str:
        return "\n".join(x for x in (self.title, self.body, self.tables) if x)


def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _frame_lines(tf) -> list[str]:
    lines = []
    for p in tf.paragraphs:
        t = "".join(r.text for r in p.runs).strip() or p.text.strip()
        if t:
            lines.append(t)
    return lines


def parse_pptx(path: Path) -> list[ParsedSlide]:
    prs = Presentation(str(path))
    out: list[ParsedSlide] = []
    for idx, slide in enumerate(prs.slides, start=1):
        ps = ParsedSlide(n=idx)
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame:
            ps.title = " ".join(_frame_lines(title_shape.text_frame)).strip()
        table_rows: list[str] = []
        for sh in _walk(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(cells):
                        table_rows.append(" | ".join(cells))
            elif sh.has_text_frame:
                ps.body_lines.extend(_frame_lines(sh.text_frame))
        if not ps.title and ps.body_lines:
            ps.title = ps.body_lines.pop(0)
        if not ps.title:
            ps.title = f"Slide {idx}"
        ps.body = "\n".join(ps.body_lines)
        ps.tables = "\n".join(table_rows)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            ps.notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append(ps)
    return out
