"""Generate tests/sample_deck.pptx (5 slides, speaker notes, one table). Usage: python scripts/make_sample_deck.py [out]"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

SLIDES = [
    ("Nimbus Analytics: 2026 Product Vision",
     ["Real-time analytics for small retail teams", "Q3 2026 all-hands"],
     "I want to introduce Nimbus Analytics, our real-time analytics product built for small retail teams. "
     "Our mission is to give a corner store the same insight a national chain has. Today I will cover the problem, "
     "the product, our traction, pricing and the roadmap."),
    ("The Problem",
     ["Small retailers lose 8 percent of revenue to stockouts", "Existing tools cost over 2,000 dollars per month",
      "Most owners decide from gut feel"],
     "Small retailers lose about eight percent of their revenue to stockouts because they find out too late. "
     "The tools that exist cost more than two thousand dollars a month, which is out of reach for a single store. "
     "So most owners end up deciding from gut feel instead of data."),
    ("The Product",
     ["Live dashboard updated every 5 seconds", "Plugs into Shopify and Square in 10 minutes",
      "Automatic reorder alerts"],
     "Nimbus connects to Shopify or Square in about ten minutes, with no engineers needed. "
     "The dashboard refreshes every five seconds, so you always see what is selling right now. "
     "The feature owners love most is automatic reorder alerts, which warn you before a shelf goes empty."),
    ("Traction",
     ["Pilot results across 3 customer tiers"],
     "We launched a pilot in March and the numbers have been encouraging. "
     "Our fastest-growing tier is the Growth plan, and weekly active usage is above ninety percent across all tiers. "
     "Customers in the pilot cut stockouts by roughly a third within two months."),
    ("Pricing and Roadmap",
     ["Starter 49 dollars, Growth 149 dollars, Scale 399 dollars", "Mobile app ships in Q4", "Multi-store rollups in Q1 2027"],
     "Pricing is simple. Starter is forty-nine dollars a month, Growth is one hundred forty-nine, and Scale is three hundred ninety-nine. "
     "On the roadmap, the mobile app ships in the fourth quarter, and multi-store rollups arrive in the first quarter of 2027. "
     "I would love to hear your questions on any of this."),
]
TABLE = [("Tier", "Customers", "MRR (USD)"), ("Starter", "42", "2,058"), ("Growth", "17", "2,533"), ("Scale", "4", "1,596")]


def build(out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, (title, bullets, notes) in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1] if i > 1 else prs.slide_layouts[0])
        slide.shapes.title.text = title
        body = slide.placeholders[1]
        body.text_frame.text = bullets[0]
        for b in bullets[1:]:
            body.text_frame.add_paragraph().text = b
        for p in body.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(24)
        if i == 4:
            rows, cols = len(TABLE), len(TABLE[0])
            tbl = slide.shapes.add_table(rows, cols, Inches(1), Inches(3.6), Inches(8), Inches(2.4)).table
            for r, row in enumerate(TABLE):
                for c, v in enumerate(row):
                    tbl.cell(r, c).text = v
        slide.notes_slide.notes_text_frame.text = notes
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "tests" / "sample_deck.pptx"
    print(build(target))
